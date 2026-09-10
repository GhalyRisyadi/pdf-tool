"""Script to generate test fixtures for pdftool test suite.
"""
from pathlib import Path
import fitz  # PyMuPDF
from PIL import Image, ExifTags
import docx
import pikepdf


FIXTURES_DIR = Path(__file__).parent / "fixtures"


def create_all_fixtures():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Sample PDF (multi-page text with metadata)
    doc = fitz.open()
    for i in range(1, 4):
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 72), f"This is page {i} of sample PDF.\nHello world from pdftool testing suite!", fontsize=14)
    doc.set_metadata({
        "title": "Sample Document Title",
        "author": "Antigravity Tester",
        "subject": "Testing Subject",
        "creator": "pytest",
        "producer": "pdftool",
    })
    sample_pdf_path = FIXTURES_DIR / "sample.pdf"
    doc.save(str(sample_pdf_path))
    doc.close()

    # 2. Sample PDF with image
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 50), "Page with an embedded image", fontsize=12)
    # Create small in-memory image
    img = Image.new("RGB", (100, 100), color=(255, 0, 0))
    img_byte_arr = fitz.io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_bytes = img_byte_arr.getvalue()
    page.insert_image(fitz.Rect(50, 100, 200, 250), stream=img_bytes)
    doc.save(str(FIXTURES_DIR / "sample_with_images.pdf"))
    doc.close()

    # 3. Sample PDF with external link
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 50), "Page with hyperlink: https://example.com", fontsize=12)
    link_rect = fitz.Rect(50, 40, 300, 70)
    page.insert_link({
        "kind": fitz.LINK_URI,
        "from": link_rect,
        "uri": "https://example.com/pdftool"
    })
    doc.save(str(FIXTURES_DIR / "sample_with_links.pdf"))
    doc.close()

    # 4. Sample PDF with table lines / text for pdfplumber
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 50), "Sample Table Page", fontsize=14)
    # Draw table border lines
    page.draw_rect(fitz.Rect(50, 100, 350, 200), width=1)
    page.draw_line(fitz.Point(50, 150), fitz.Point(350, 150), width=1)
    page.draw_line(fitz.Point(200, 100), fitz.Point(200, 200), width=1)
    page.insert_text((60, 130), "Header 1", fontsize=11)
    page.insert_text((210, 130), "Header 2", fontsize=11)
    page.insert_text((60, 180), "Data 1", fontsize=11)
    page.insert_text((210, 180), "Data 2", fontsize=11)
    doc.save(str(FIXTURES_DIR / "sample_with_tables.pdf"))
    doc.close()

    # 5. Sample encrypted PDF (password: "secret123")
    with pikepdf.open(str(sample_pdf_path)) as pdf:
        pdf.save(
            str(FIXTURES_DIR / "sample_encrypted.pdf"),
            encryption=pikepdf.Encryption(
                owner="admin123",
                user="secret123",
                R=6,
            )
        )

    # 6. Sample restricted PDF (owner password only, empty user password)
    perms = pikepdf.Permissions(
        extract=False,
        modify_annotation=False,
        modify_assembly=False,
        modify_form=False,
        modify_other=False,
        print_lowres=False,
        print_highres=False,
    )
    with pikepdf.open(str(sample_pdf_path)) as pdf:
        pdf.save(
            str(FIXTURES_DIR / "sample_restricted.pdf"),
            encryption=pikepdf.Encryption(
                owner="admin123",
                user="",
                allow=perms,
                R=6,
            )
        )

    # 7. Sample PDF with active content (JavaScript & OpenAction) for sanitize testing
    with pikepdf.open(str(sample_pdf_path)) as pdf:
        pdf.Root.OpenAction = pikepdf.Dictionary(
            S=pikepdf.Name.JavaScript,
            JS=pikepdf.String("app.alert('Hello');")
        )
        pdf.Root.Names = pikepdf.Dictionary(
            JavaScript=pikepdf.Dictionary(
                Names=pikepdf.Array(["sample_script", pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String("1+1;"))])
            )
        )
        pdf.save(str(FIXTURES_DIR / "sample_active.pdf"))

    # 8. Sample PDF containing PII
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    pii_text = (
        "Personal Information Document:\n"
        "Contact Email: user.testing@example.com\n"
        "Phone Number: 081234567890\n"
        "Server IP: 192.168.1.100\n"
        "NIK: 3171010101900001\n"
        "Credit Card: 4532015012345675\n"
    )
    page.insert_text((50, 72), pii_text, fontsize=12)
    doc.save(str(FIXTURES_DIR / "sample_pii.pdf"))
    doc.close()

    # 9. Sample JPEG with EXIF metadata
    img = Image.new("RGB", (200, 200), color=(73, 109, 137))
    exif = img.getexif()
    exif[ExifTags.Base.Make] = "Nikon"
    exif[ExifTags.Base.Model] = "D850"
    exif[ExifTags.Base.Orientation] = 1
    # GPS IFD
    gps_ifd = {
        ExifTags.GPS.GPSLatitude: (6.0, 12.0, 0.0),
        ExifTags.GPS.GPSLatitudeRef: "S",
        ExifTags.GPS.GPSLongitude: (106.0, 48.0, 0.0),
        ExifTags.GPS.GPSLongitudeRef: "E",
    }
    exif[ExifTags.IFD.GPSInfo] = gps_ifd
    sample_jpg_path = FIXTURES_DIR / "sample.jpg"
    img.save(str(sample_jpg_path), "JPEG", exif=exif)

    # 10. Sample PNG
    png_img = Image.new("RGBA", (150, 150), color=(100, 200, 100, 255))
    png_img.save(str(FIXTURES_DIR / "sample.png"), "PNG")

    # 11. Sample DOCX
    docx_doc = docx.Document()
    docx_doc.core_properties.title = "Sample DOCX Title"
    docx_doc.core_properties.author = "Antigravity"
    docx_doc.add_heading("Sample Document Heading", 0)
    docx_doc.add_paragraph("This is the first paragraph with some sample text.")
    docx_doc.add_page_break()
    docx_doc.add_paragraph("This is the second paragraph on page two after a page break.")
    docx_doc.save(str(FIXTURES_DIR / "sample.docx"))

    print(f"[✓] Successfully generated fixtures in {FIXTURES_DIR}")


if __name__ == "__main__":
    create_all_fixtures()
