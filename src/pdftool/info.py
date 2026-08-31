from pathlib import Path
from .utils import has_javascript, get_attachments, count_images, get_fonts


def jpg_analysis(path: Path):
    from PIL import Image, ExifTags

    try:

        size_kb = path.stat().st_size / 1024
        size_mb = size_kb / 1024

        with Image.open (path) as img:
            exif = img.getexif()

            print("JPG Analysis")
            print("─" * 50)
            print(f"File    : {path.name}")
            print(f"Size    : {size_kb:.1f} KB  ({size_mb:.2f} MB)")
            print(f"Dimensions : {img.width} x {img.height} px")
            print(f"Format  : {img.format}")
            print(f"Mode    : {img.mode}")

            print("\nEXIF:")

            if not exif:
                print("  None")
                return

            make = exif.get(ExifTags.Base.Make)
            model = exif.get(ExifTags.Base.Model)
            orientation = exif.get(ExifTags.Base.Orientation)
            exif_ifd = exif.get_ifd(ExifTags.IFD.Exif)
            date_taken = exif_ifd.get(ExifTags.Base.DateTimeOriginal)
            gps_ifd = exif.get_ifd(ExifTags.IFD.GPSInfo)

            print(f"  Camera     : {make or '-'}")
            print(f"  Model      : {model or '-'}")
            print(f"  Date Taken : {date_taken or '-'}")
            print(f"  Orientation: {orientation or '-'}")

            if gps_ifd:
                    gps_latitude = gps_ifd.get(ExifTags.GPS.GPSLatitude)
                    gps_latitude_ref = gps_ifd.get(ExifTags.GPS.GPSLatitudeRef)

                    gps_longitude = gps_ifd.get(ExifTags.GPS.GPSLongitude)
                    gps_longitude_ref = gps_ifd.get(ExifTags.GPS.GPSLongitudeRef)

                    if (
                        gps_latitude is not None
                        and gps_latitude_ref is not None
                        and gps_longitude is not None
                        and gps_longitude_ref is not None
                    ):

                        def format_gps(coords, ref):
                            try:
                                d = float(coords[0])
                                m = float(coords[1])
                                s = float(coords[2])

                                dec = d + (m / 60.0) + (s / 3600.0)

                                ref_str = str(ref).strip().upper()
                                if ref_str in ['S', 'W']:
                                    dec = -dec

                                dms = f"{int(d)}° {int(m)}' {s:.3f}\" {ref_str}"
                                return dms, dec
                            except Exception:
                                return f"{coords} {ref}", None
                            
                        lat_dms, lat_dec = format_gps(gps_latitude, gps_latitude_ref)
                        lon_dms, lon_dec = format_gps(gps_longitude, gps_longitude_ref)

                        print("\nGPS:")
                        print(f"  Latitude   : {lat_dms}")
                        print(f"  Longitude  : {lon_dms}")

                        if lat_dec is not None and lon_dec is not None:
                            print(f"  Coordinates : {lat_dec:.6f}, {lon_dec:.6f}")
                            print(f"  Map         : https://www.google.com/maps/search/?api=1&query={lat_dec:.6f},{lon_dec:.6f}")

                    else:
                        print("\nGPS")
                        print("  Present, coordinates unavailable")
            else:
                print("\nGPS")
                print("  None")

    except Exception as e:
        print(f"\n [ERROR] Failed to read JPG info: {e}")

def doc_info(path: Path, origin: Path | None = None):
    from docx import Document

    origin = origin or path

    try:
        doc    = Document(str(path))
        meta   = doc.core_properties
        size   = path.stat().st_size / 1024
        n_word = sum(len(p.text.split()) for p in doc.paragraphs)
    
        print(f"File       : {path.name}")
        print(f"Size       : {size:.1f} KB  ({size / 1024:.2f} MB)")
        print(f"Paragraphs : {len(doc.paragraphs)}")
        print(f"Words      : {n_word}")
        if meta.title:  print(f"Title      : {meta.title}")
        if meta.author: print(f"Author     : {meta.author}")

    except Exception as e:
        print(f"\n[ERROR] Failed to read DOC info: {e}")

def pdf_analysis(path: Path):
    from pypdf import PdfReader

    try:
        reader  = PdfReader(str(path))
        size_kb = path.stat().st_size / 1024
        size_mb = size_kb / 1024
        
        is_encrypted = reader.is_encrypted
        needs_password = False

        if is_encrypted:
            if reader.decrypt("") == 0:
                needs_password = True
        
        print("PDF Analysis")
        print("─" * 50)
        print(f"File        : {path.name}")
        print(f"Size        : {size_kb:.1f} KB ({size_mb:.2f} MB)")

        if needs_password:
            javascript = None
            fonts = set()
            print(f"Pages       : (need password to read)")
            print(f"Attachments : (need password to read)")
            print(f"Images      : (need password to read)")
            print(f"Fonts       : (need password to read)")
        else:
            meta = reader.metadata
            javascript = has_javascript(reader)
            attachments = get_attachments(reader)
            fonts = get_fonts(reader)
            images = count_images(reader)

            print(f"Pages       : {len(reader.pages)}")
            print(f"Attachments : {len(attachments)}")
            print(f"Images      : {images}")
            print(f"Fonts       : {len(fonts)}")

        print("\nMetadata:")
        if needs_password:
            print("  (need password to read metadata)")
        elif meta:
            print(f"  Title     : {meta.title or '-'}")
            print(f"  Author    : {meta.author or '-'}")
            print(f"  Subject   : {meta.subject or '-'}")
            print(f"  Creator   : {meta.creator or '-'}")
            print(f"  Producer  : {meta.producer or '-'}")
        else:
            print("  None")

        print("\nSecurity")
        _print_security_info(path, is_encrypted, javascript)

        if fonts:
            print(f"\nFonts: {len(fonts)}")

            for font in sorted(fonts):
                print(f"  • {font}")

    except Exception as e:
        print(f"\n[ERROR] Failed to read PDF file: {e}")

_ENC_METHOD_LABELS = {
    "none": "RC4",
    "rc4":  "RC4",
    "aes":  "AES",
    "aesv3": "AES",
}

def _print_security_info(path: Path, is_encrypted: bool, javascript: bool | None):
    
    if not is_encrypted:
        print(f"  Encrypted   : No")
        print(f"  JavaScript  : {'Yes' if javascript else 'No'}")
        print(f"  Permissions : -")
        return
 
    import pikepdf
 
    try:
        with pikepdf.open(str(path)) as pdf:
            enc = pdf.encryption
            allow = pdf.allow
 
            method = _ENC_METHOD_LABELS.get(str(enc.stream_method).split(".")[-1], "Unknown")
            algo = f"{method}-{enc.bits}"
 
            if allow.print_highres:
                print_status = "Allowed"
            elif allow.print_lowres:
                print_status = "Low-res only"
            else:
                print_status = "Denied"
 
            js_label = "(not yet checked)" if javascript is None else ("Yes" if javascript else "No")
 
            print(f"  Encrypted   : Yes")
            print(f"  Encryption  : {algo}")
            print(f"  JavaScript  : {js_label}")
            print(f"  Print       : {print_status}")
            print(f"  Copy        : {'Allowed' if allow.extract else 'Denied'}")
            print(f"  Modify      : {'Allowed' if allow.modify_other else 'Denied'}")
 
    except pikepdf.PasswordError:
        print(f"  Encrypted   : Yes")
        print(f"  Encryption  : (need password to read)")
        print(f"  JavaScript  : (need password to read)")
        print(f"  Permissions : (need password to read)")